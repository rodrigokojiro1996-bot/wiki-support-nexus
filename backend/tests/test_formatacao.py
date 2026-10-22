"""O importador gera os mesmos blocos que o editor guarda para o texto colado (exemplos em dados/importacao/formatacao.json).

O outro lado (como o editor mostra esses blocos) é testado em frontend/src/componentes/textoRico.test.tsx.
Aqui também ficam os estilos do Word: títulos, listas numeradas e com marcador e fonte monoespaçada.
"""
import io
import json
from pathlib import Path

import docx
import pytest
from docx.shared import Pt

from app.servicos.importacao import converter

CASOS = json.loads((Path(__file__).parent / "dados" / "importacao" / "formatacao.json").read_text(encoding="utf-8"))["casos"]


@pytest.mark.parametrize("caso", CASOS, ids=[c["nome"] for c in CASOS])
def test_texto_do_documento_gera_os_blocos_esperados(banco, caso):
    r = converter(caso["texto"].encode(), "exemplo.txt", {"caixa_mov"}, secao_unica=None if "secoes" in caso else "Trecho")
    if "secoes" in caso:
        assert r["secoes"] == caso["secoes"]
    else:
        assert r["secoes"] == [{"titulo": "Trecho", "blocos": caso["blocos"]}]


def test_casos_com_mesma_aparencia_sao_um_bloco_de_texto():
    # o editor guarda o texto colado num bloco de texto só; é essa a comparação feita no frontend
    for c in CASOS:
        if c.get("mesma_aparencia"):
            assert [b["tipo"] for b in c["blocos"]] == ["texto"], c["nome"]


def criar_docx_com_estilos() -> bytes:
    d = docx.Document()
    d.add_heading("Fechamento do caixa", level=1)
    d.add_heading("Antes de começar", level=2)
    d.add_paragraph("Confira os itens:")
    d.add_paragraph("Sangrias lançadas", style="List Bullet")
    d.add_paragraph("Cartões conferidos", style="List Bullet 2")
    d.add_paragraph("Acesse Caixa > Fechamento.", style="List Number")
    d.add_paragraph("Informe o valor contado.", style="List Number")
    for linha in ("SELECT total", "FROM caixa_mov;"):
        run = d.add_paragraph().add_run(linha)
        run.font.name, run.font.size = "Consolas", Pt(10)
    d.add_paragraph("Texto comum depois do código.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_word_usa_os_estilos_do_documento(banco):
    r = converter(criar_docx_com_estilos(), "caixa.docx", {"caixa_mov"})
    assert r["secoes"] == [{"titulo": "Fechamento do caixa", "blocos": [
        {"tipo": "subtitulo", "texto": "Antes de começar"},                                   # Título 2
        {"tipo": "texto", "texto": "Confira os itens:\n• Sangrias lançadas\n\t• Cartões conferidos"},  # marcadores
        {"tipo": "passo", "rotulo": "1", "texto": "Acesse Caixa > Fechamento."},                # lista numerada
        {"tipo": "passo", "rotulo": "2", "texto": "Informe o valor contado."},
        {"tipo": "sql", "texto": "SELECT total\nFROM caixa_mov;"},                               # fonte monoespaçada
        {"tipo": "texto", "texto": "Texto comum depois do código."},
    ]}]
    assert r["tabelas"] == ["caixa_mov"]
