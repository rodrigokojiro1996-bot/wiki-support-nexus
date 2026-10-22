"""Importação de documentos em rascunho de manual (app/servicos/importacao.py e POST /api/manuais/importar).

Só documentos fictícios: os de tests/dados/importacao/ (gerados por gerar_exemplos.py) e os montados aqui mesmo.
"""
import io
import json
import random
import zipfile
from pathlib import Path
from urllib.parse import quote

import docx
import pytest
from PIL import Image

from app.servicos import importacao
from app.servicos.importacao import ImportacaoRecusada, converter
from ferramentas import importar_processos as lote

EXEMPLOS = Path(__file__).parent / "dados" / "importacao"
TABELAS = {"caixa_mov", "formas_pgto", "transferencias"}
OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600  # cabeçalho de um .doc antigo


def blocos(r: dict, tipo: str) -> list[dict]:
    return [b for s in r["secoes"] for b in s["blocos"] if b["tipo"] == tipo]


def png(largura: int, altura: int, semente: int) -> bytes:
    """PNG com ruído: comprime pouco e passa do tamanho mínimo de print."""
    rnd = random.Random(semente)
    img = Image.new("RGB", (largura, altura))
    img.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)) for _ in range(largura * altura)])
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


# ---------------- texto ----------------
def test_txt_vira_secoes_passos_e_sql(banco):
    r = converter((EXEMPLOS / "exemplo.txt").read_bytes(), "exemplo.txt", TABELAS)
    assert [s["titulo"] for s in r["secoes"]][1:] == ["Transferência presa como em uso", "Transferência não aparece no destino"]
    assert [b["rotulo"] for b in blocos(r, "passo")] == ["1.1", "1.2"]
    assert blocos(r, "sql")[0]["texto"].startswith("SELECT id_transf")
    assert r["tabelas"] == ["transferencias"]
    assert r["fontes"] == ["exemplo.txt"]


def test_senha_sai_mascarada(banco):
    r = converter((EXEMPLOS / "exemplo.txt").read_bytes(), "exemplo.txt", TABELAS)
    texto = json.dumps(r, ensure_ascii=False)
    assert "segredo123" not in texto
    assert "password=******" in texto
    assert "1 senha mascarada (******): confira se o texto continua claro." in r["avisos"]


def test_aviso_conta_cada_senha_uma_vez(banco):
    r = converter("Acesso\npassword=abc e --password=xyz\n".encode(), "acesso.txt", set())
    assert "2 senhas mascaradas (******): confira se o texto continua claro." in r["avisos"]


def test_txt_utf8_com_bom_e_cp1252_geram_os_mesmos_blocos(banco):
    texto = (EXEMPLOS / "exemplo.txt").read_text(encoding="utf-8")
    utf8 = converter(texto.encode("utf-8"), "exemplo.txt", TABELAS)
    bom = converter(texto.encode("utf-8-sig"), "exemplo.txt", TABELAS)
    cp1252 = converter(texto.encode("cp1252"), "exemplo.txt", TABELAS)
    assert utf8["secoes"] == bom["secoes"] == cp1252["secoes"]
    assert not any("cp1252" in a for a in utf8["avisos"] + bom["avisos"])
    assert any("cp1252" in a for a in cp1252["avisos"])


def test_sem_titulo_avisa_e_usa_o_nome_do_arquivo(banco):
    r = converter("Confira o cadastro.\nDepois feche a tela.\n".encode(), "conferir_cadastro.txt", set())
    assert r["titulo"] == "Conferir cadastro"
    assert "Nenhum título encontrado: todo o conteúdo ficou em uma seção." in r["avisos"]


def test_md_aceita_titulos_e_bloco_sql(banco):
    md = "# Consultar vendas\nVeja o total do dia.\n```sql\nSELECT total FROM caixa_mov;\n```\n"
    r = converter(md.encode(), "vendas.md", TABELAS)
    assert r["formato"] == "md"
    assert r["titulo"] == "Consultar vendas"
    assert blocos(r, "sql") == [{"tipo": "sql", "texto": "SELECT total FROM caixa_mov;"}]


# ---------------- PDF ----------------
def test_pdf_json_esperado(banco):
    r = converter((EXEMPLOS / "exemplo.pdf").read_bytes(), "exemplo.pdf", TABELAS)
    assert "Manual de exemplo" not in json.dumps(r, ensure_ascii=False)  # cabeçalho repetido nas páginas
    imagens = blocos(r, "imagem")
    assert len(imagens) == 1  # o logo aparece nas 3 páginas e é descartado
    assert imagens[0]["arquivo"].startswith("_rascunho/")
    assert (banco / "imagens" / imagens[0]["arquivo"]).exists()
    imagens[0]["arquivo"] = "IMG"
    assert r["titulo"] == "Abrir o caixa"
    assert [s["titulo"] for s in r["secoes"]] == ["Abrir o caixa", "Conferir o fechamento"]
    assert r["secoes"][0]["blocos"] == [
        {"tipo": "passo", "rotulo": "1.1", "texto": "Acesse Caixa > Abertura de Caixa."},
        {"tipo": "passo", "rotulo": "1.2", "texto": "Informe o valor inicial e confirme."},
        {"tipo": "obs", "texto": "o caixa só abre com o operador logado."},
        {"tipo": "imagem", "arquivo": "IMG", "legenda": ""},
    ]
    assert r["secoes"][1]["blocos"][:3] == [
        {"tipo": "passo", "rotulo": "2.1", "texto": "Acesse Caixa > Fechamento de Caixa."},
        {"tipo": "nota", "texto": "confira as sangrias antes de fechar."},
        {"tipo": "sql", "texto": "SELECT * FROM caixa_mov WHERE id_caixa = 10;"},  # o "*" fica
    ]
    assert r["tabelas"] == ["caixa_mov"]
    assert "3 imagens ignoradas por se repetir no documento (logo, rodapé)." in r["avisos"]
    assert "1 senha mascarada (******): confira se o texto continua claro." in r["avisos"]


def test_pdf_acima_do_limite_de_paginas(banco, monkeypatch):
    monkeypatch.setattr(importacao, "PAGINAS_MAXIMO", 2)
    with pytest.raises(ImportacaoRecusada) as e:
        converter((EXEMPLOS / "exemplo.pdf").read_bytes(), "exemplo.pdf", TABELAS)
    assert e.value.status == 413 and "3 páginas" in e.value.mensagem


def test_pdf_acima_do_limite_de_imagens(banco, monkeypatch):
    monkeypatch.setattr(importacao, "IMAGENS_MAXIMO", 2)
    r = converter((EXEMPLOS / "exemplo.pdf").read_bytes(), "exemplo.pdf", TABELAS)
    assert "O documento tem mais de 2 imagens: só as primeiras foram importadas." in r["avisos"]


# ---------------- Word ----------------
def criar_docx() -> bytes:
    """Documento Word fictício: títulos, passo, OBS, SQL, um print, um logo repetido 3 vezes e uma tabela."""
    d = docx.Document()
    logo = png(120, 45, 1)
    d.add_picture(io.BytesIO(logo))
    d.add_heading("Abrir o caixa", level=1)
    d.add_paragraph("1.1 Acesse Caixa > Abertura de Caixa.")
    d.add_paragraph("OBS: o caixa só abre com o operador logado.")
    d.add_paragraph("SELECT id_caixa FROM caixa_mov;")
    d.add_picture(io.BytesIO(png(200, 100, 2)))
    d.add_picture(io.BytesIO(logo))
    d.add_heading("Formas de pagamento", level=1)
    t = d.add_table(rows=2, cols=2)
    for i, (a, b) in enumerate([("Forma", "Código"), ("Dinheiro", "01")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    d.add_picture(io.BytesIO(logo))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_json_esperado(banco):
    r = converter(criar_docx(), "caixa.docx", TABELAS)
    assert r["formato"] == "docx"
    imagens = blocos(r, "imagem")
    assert len(imagens) == 1 and imagens[0]["arquivo"].startswith("_rascunho/")
    imagens[0]["arquivo"] = "IMG"
    assert r["secoes"] == [
        {"titulo": "Abrir o caixa", "blocos": [
            {"tipo": "passo", "rotulo": "1.1", "texto": "Acesse Caixa > Abertura de Caixa."},
            {"tipo": "obs", "texto": "o caixa só abre com o operador logado."},
            {"tipo": "sql", "texto": "SELECT id_caixa FROM caixa_mov;"},
            {"tipo": "imagem", "arquivo": "IMG", "legenda": ""},
        ]},
        {"titulo": "Formas de pagamento", "blocos": [
            {"tipo": "texto", "texto": "Forma | Código"},
            {"tipo": "texto", "texto": "Dinheiro | 01"},
        ]},
    ]
    assert r["titulo"] == "Abrir o caixa"
    assert r["tabelas"] == ["caixa_mov"]
    assert "3 imagens ignoradas por se repetir no documento (logo, rodapé)." in r["avisos"]
    assert importacao.AVISO_TABELA in r["avisos"]


# ---------------- tipo do arquivo ----------------
def zip_com(arquivos: dict[str, bytes], compressao: int = zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compressao) as z:
        for nome, conteudo in arquivos.items():
            z.writestr(nome, conteudo)
    return buf.getvalue()


@pytest.mark.parametrize("dados, nome, status, trecho", [
    (OLE, "antigo.doc", 415, "salve como .docx"),
    (OLE, "renomeado.docx", 415, "salve como .docx"),  # vale o conteúdo, não a extensão
    ("texto qualquer".encode(), "planilha.doc", 415, "salve como .docx"),
    (b"PK\x03\x04", "macro.docm", 415, "sem macros"),
    (zip_com({"word/document.xml": b"<w/>", "word/vbaProject.bin": b"x"}), "macro.docx", 415, "sem macros"),
    (zip_com({"xl/workbook.xml": b"<x/>"}), "planilha.docx", 415, "não é um documento do Word"),
    (zip_com({"word/document.xml": b"\0" * 5_000_000}), "bomba.docx", 413, "grande demais"),  # taxa ~1000:1
    (png(100, 50, 3), "print.pdf", 415, "não suportado"),
    (b"   \n ", "vazio.txt", 422, "vazio"),
    (b"%PDF-1.7\nlixo", "quebrado.pdf", 422, "danificado"),
], ids=lambda v: v if isinstance(v, str) else "-")
def test_arquivos_recusados(banco, dados, nome, status, trecho):
    with pytest.raises(ImportacaoRecusada) as e:
        converter(dados, nome, set())
    assert e.value.status == status
    assert trecho in e.value.mensagem


def test_docx_grande_demais_descompactado(banco, monkeypatch):
    monkeypatch.setattr(importacao, "DESCOMPACTADO_MAXIMO", 1000)
    sem_compressao = zip_com({"word/document.xml": b"<w/>" * 500}, zipfile.ZIP_STORED)  # taxa 1:1, só o total estoura
    with pytest.raises(ImportacaoRecusada) as e:
        converter(sem_compressao, "grande.docx", set())
    assert e.value.status == 413


# ---------------- rota ----------------
def importar(cliente, dados: bytes, nome: str, modulo: str = "pdv"):
    return cliente.post(f"/api/manuais/importar?modulo={modulo}", content=dados, headers={"X-Arquivo": quote(nome)})


def test_rota_devolve_rascunho_sem_salvar(cliente, banco):
    antes = cliente.get("/api/manuais").json()
    r = importar(cliente, (EXEMPLOS / "exemplo.txt").read_bytes(), "C:\\pasta\\Transferência entre lojas.txt")
    assert r.status_code == 200
    corpo = r.json()
    assert set(corpo) >= {"titulo", "subtitulo", "secoes", "fontes", "tabelas", "avisos"}
    assert corpo["fontes"] == ["Transferência entre lojas.txt"]  # só o nome, sem a pasta
    assert corpo["tabelas"] == ["transferencias"]  # do banco do módulo (demo_pdv)
    assert cliente.get("/api/manuais").json() == antes
    assert not list((banco / "manuais").glob("*.json"))


def test_rota_413_415_e_404(cliente, monkeypatch):
    with monkeypatch.context() as m:
        m.setattr(importacao, "TAMANHO_MAXIMO", 1000)
        r = importar(cliente, b"a" * 1001, "grande.txt")
        assert r.status_code == 413 and "MB" in r.json()["detail"]
    r = importar(cliente, png(100, 50, 3), "print.png")
    assert r.status_code == 415 and "Envie PDF" in r.json()["detail"]
    r = importar(cliente, OLE, "antigo.doc")
    assert r.status_code == 415 and ".docx" in r.json()["detail"]
    r = importar(cliente, b"PK\x03\x04", "macro.docm")
    assert r.status_code == 415 and "macros" in r.json()["detail"]
    r = importar(cliente, zip_com({"word/document.xml": b"\0" * 5_000_000}), "bomba.docx")
    assert r.status_code == 413
    assert importar(cliente, b"texto", "a.txt", modulo="xyz").status_code == 404


def test_rota_conta_importacoes_nas_metricas(cliente):
    importar(cliente, (EXEMPLOS / "exemplo.txt").read_bytes(), "exemplo.txt")
    importar(cliente, png(100, 50, 3), "print.png")
    m = cliente.get("/metrics").text
    assert 'nexus_importacoes_total{formato="txt",resultado="ok"}' in m
    assert 'nexus_importacoes_total{formato="desconhecido",resultado="recusado"}' in m
    assert 'nexus_importacao_duration_seconds_count{formato="txt"}' in m


def test_fluxo_completo_importar_revisar_e_salvar(cliente, banco):
    rascunho = importar(cliente, (EXEMPLOS / "exemplo.pdf").read_bytes(), "exemplo.pdf").json()
    print_rascunho = blocos(rascunho, "imagem")[0]["arquivo"]
    corpo = {"titulo": rascunho["titulo"], "subtitulo": "Do PDF de exemplo", "modulo": "pdv",
             "secoes": rascunho["secoes"], "fontes": rascunho["fontes"]}
    r = cliente.post("/api/manuais", json=corpo, headers={"X-Usuario": "Ana"})
    assert r.status_code == 201, r.text
    codigo = r.json()["codigo"]
    m = cliente.get(f"/api/manuais/{codigo}").json()
    assert m["fontes"] == ["exemplo.pdf"]
    assert m["tabelas"] == ["caixa_mov"]
    img = blocos(m, "imagem")[0]["arquivo"]
    assert img.startswith(f"{codigo}/") and (banco / "imagens" / img).exists()
    assert not (banco / "imagens" / print_rascunho).exists()  # o print saiu do rascunho
    # editar no editor comum (sem mandar fontes) mantém as do documento
    corpo.pop("fontes")
    corpo["secoes"] = m["secoes"]
    assert cliente.put(f"/api/manuais/{codigo}", json=corpo, headers={"X-Usuario": "Ana"}).status_code == 200
    assert cliente.get(f"/api/manuais/{codigo}").json()["fontes"] == ["exemplo.pdf"]


# ---------------- lote ----------------
def test_lote_grava_o_mesmo_conteudo_da_conversao(banco, tmp_path):
    origem = tmp_path / "docs"
    origem.mkdir()
    for n in ("exemplo.pdf", "exemplo.txt"):
        (origem / n).write_bytes((EXEMPLOS / n).read_bytes())
    cur = tmp_path / "curadoria.json"
    cur.write_text(json.dumps({"origem": str(origem), "manuais": [
        {"codigo": "T.PDF", "titulo": "Caixa", "menu": "pdv.caixa.fechamento-de-caixa", "fontes": [{"arquivo": "exemplo.pdf"}]},
        {"codigo": "T.TXT", "titulo": "Transferências", "menu": "pdv.caixa.fechamento-de-caixa", "fontes": [{"arquivo": "exemplo.txt"}]},
    ]}), encoding="utf-8")
    lote.main(cur)
    for codigo, n in (("T.PDF", "exemplo.pdf"), ("T.TXT", "exemplo.txt")):
        gravado = json.loads((banco / "manuais" / f"{codigo}.json").read_text(encoding="utf-8"))
        r = converter((EXEMPLOS / n).read_bytes(), n, TABELAS)
        # no lote o print vai direto para imagens/<código>/; na interface, para o rascunho
        for b_lote, b_conv in zip(blocos(gravado, "imagem"), blocos(r, "imagem"), strict=True):
            assert b_lote["arquivo"] == f"{codigo}/01.png"
            assert (banco / "imagens" / b_lote["arquivo"]).read_bytes() == (banco / "imagens" / b_conv["arquivo"]).read_bytes()
            b_lote["arquivo"] = b_conv["arquivo"]
        assert gravado["secoes"] == r["secoes"]
        assert gravado["tabelas"] == r["tabelas"]


def test_lote_lista_documento_recusado_como_problema(banco, tmp_path):
    origem = tmp_path / "docs"
    origem.mkdir()
    (origem / "antigo.doc").write_bytes(OLE)
    cur = tmp_path / "curadoria.json"
    cur.write_text(json.dumps({"origem": str(origem), "manuais": [
        {"codigo": "T.DOC", "titulo": "Antigo", "menu": "pdv.caixa.fechamento-de-caixa", "fontes": [{"arquivo": "antigo.doc"}]},
    ]}), encoding="utf-8")
    problemas = lote.main(cur)
    assert any("antigo.doc" in p and ".docx" in p for p in problemas)
